# -*- coding: utf-8 -*-
"""离线自测：截图只导出去重后的唯一发票（2026-10-08 用户需求）

需求原文：
「批量导入时已识别到重复发票，汇总页显示了重复，输出的查验截图就应该只导出
 唯一的一张不重复的截图（以导入文件中第一张不重复的发票文件名命名），
 这样我就不用再手动去重了。」

覆盖 5 件事：
  1) 35 条记录 / 7 张唯一发票 → 截图目录只落地 7 张
  2) 文件名 = 每组「第一张」发票的文件名（不是副本的）
  3) HTML 不重复出图（<img> = 7）、下载按钮张数 = 7
  4) Excel 重复行的「查验截图文件」列标注「共用同一张截图」
  5) prune_dup_shots 只移走逐字节相同的副本；
     内容不同的同名文件绝不动（防止误删）
"""
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import extract                                    # noqa: E402
import report                                     # noqa: E402

# 产物写系统临时目录，绝不污染项目目录
OUT = Path(tempfile.mkdtemp(prefix="invoice_dedupe_"))
(OUT / "raw").mkdir(parents=True)
(OUT / "查验截图").mkdir(parents=True)

BAD = 0


def check(cond, msg):
    global BAD
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        BAD += 1


# ---------- 造数：7 张唯一发票，共 35 条记录（模拟用户实际的 35 行）----------
GROUPS = [7, 6, 9, 1, 1, 10, 1]                  # 每组重复次数，合计 35
NUMS = ["25952000000000000001", "25952000000000000002", "25952000000000000003",
        "25952000000000000004", "26952000000000000006", "25952000000000000005",
        "26952000000000000007"]

from PIL import Image                             # noqa: E402
raw_shots = {}
for gi, _num in enumerate(NUMS, 1):
    p = OUT / "raw" / f"shot{gi}.png"
    Image.new("RGB", (400, 260), (200 + gi * 5, 220, 250)).save(p)
    raw_shots[gi] = p

recs, idx = [], 0
masters = {}
for gi, (num, cnt) in enumerate(zip(NUMS, GROUPS), 1):
    for k in range(cnt):
        idx += 1
        recs.append({"idx": idx, "file": f"PO-num{gi:02d}-copy{k + 1}.PDF",
                     "number": num, "date": "20250512", "amount": "53803.00",
                     "kind": "数电票", "code": "", "checkcode": "", "note": ""})

audit = extract.audit_duplicates(recs)
todo, followers = extract.plan_dedupe(recs, audit)
print(f"清单 {len(recs)} 行 → 唯一发票 {len(todo)} 张，合并重复 {len(followers)} 行")

results = []
for r in todo:
    gi = int(r["file"].split("num")[1][:2])
    rr = dict(r)
    rr["result"] = "一致"
    rr["detail"] = "当日第1次，查验通过"
    rr["shot"] = str(raw_shots[gi])
    masters[r["idx"]] = rr
    results.append(rr)
for f in followers:
    rr = dict(f["rec"])
    m = masters[f["master"]["idx"]]
    rr.update({"result": m["result"], "shot": m["shot"], "is_dup": True,
               "dup_of": f["master"]["idx"],
               "detail": f"【重复发票】与「{Path(m['file']).name}」为同一张发票"})
    results.append(rr)
results.sort(key=lambda x: x["idx"])

# ---------- 1 + 2：截图只落地唯一发票 ----------
print("\n[1/5] 截图落地数量与命名")
mapping = report.materialize_shots(results, OUT / "查验截图")
files = sorted(p.name for p in (OUT / "查验截图").glob("*.png"))
check(len(results) == 35, f"记录数 {len(results)} = 35")
check(len(mapping) == 7, f"截图落地 {len(mapping)} 张 = 7（唯一发票）")
check(len(files) == 7, f"截图目录实际文件 {len(files)} 个 = 7")
want = {f"PO-num{gi:02d}-copy1_查验截图.png" for gi in range(1, 8)}
check(set(files) == want, "命名 = 每组第一张不重复的发票文件名（copy1）")
check(all("copy1" in f for f in files),
      "目录里没有 copy2/copy3… 的重复副本")

# 重复行不应该有 shot_final
n_dup_final = sum(1 for r in results if r.get("is_dup") and r.get("shot_final"))
check(n_dup_final == 0, f"重复行没有落地截图（{n_dup_final} 个）")

# ---------- 3：HTML ----------
print("\n[2/5] HTML 报告")
zip_path = report.pack_shots_zip(mapping, str(OUT / "查验截图.zip"))
zip_ok = report.pack_shots_zip(mapping, str(OUT / "查验成功截图.zip"))
htm = report.write_html(results, str(OUT / "查验结果.html"),
                        zip_name=Path(zip_path).name,
                        zip_ok_name=Path(zip_ok).name,
                        shots_out_dir=OUT / "查验截图")
txt = Path(htm).read_text(encoding="utf-8")
n_img = txt.count("<img ")
check(n_img == 7, f"HTML 缩略图 {n_img} 张 = 7（重复行不再重复出图）")
check("zip，7 张" in txt, "下载按钮张数写的是 7 张")
check("同一张票" in txt, "重复行标注「与第 N 行同一张票」")
with zipfile.ZipFile(zip_path) as z:
    nz = len(z.namelist())
check(nz == 7, f"全部截图.zip 内 {nz} 个文件 = 7")

# ---------- 4：Excel ----------
print("\n[3/5] Excel 结果表")
xlsx = report.write_excel(results, str(OUT / "查验结果.xlsx"))
from openpyxl import load_workbook                  # noqa: E402
ws = load_workbook(xlsx, data_only=True).active
col = [r[10] for r in ws.iter_rows(min_row=2, values_only=True)]
check(len(col) == 35, f"Excel 行数 {len(col)} = 35")
check(sum(1 for v in col if v and "共用同一张截图" in str(v)) == 28,
      "28 个重复行都标注了「共用同一张截图」")
check(sum(1 for v in col if v and v.endswith("_查验截图.png")) == 7,
      "只有 7 行指向真实截图文件")

# ---------- 5：prune 只移走逐字节相同的副本 ----------
print("\n[4/5] 历史重复副本清理")
sh = OUT / "查验截图"
# 造出 28 个「上一版遗留」的重复副本 + 1 个同名但内容不同的文件
import hashlib                                       # noqa: E402
dup_rows = [r for r in results if r.get("is_dup")]
for r in dup_rows:
    src = masters[r["dup_of"]]["shot"]
    shutil.copy2(src, sh / (report.safe_stem(Path(r["file"]).stem)
                            + "_查验截图.png"))
trap = sh / "PO-num01-copy2_查验截图.png"
Image.new("RGB", (400, 260), (10, 10, 10)).save(trap)   # 内容不同，不能被误删
before = len(list(sh.glob("*.png")))
moved = report.prune_dup_shots(results, sh)
after = sorted(p.name for p in sh.glob("*.png"))
check(len(moved) == 27, f"移走 27 个真副本（1 个内容不同的被保留）：{len(moved)}")
check(before - len(after) == 27, f"目录从 {before} 降到 {len(after)}")
check(trap.name in after, "内容不同的同名文件未被误删")
check(all(p.is_file() for p in sh.glob("*.png")), "目录仍可读")

# ---------- 归档落点 ----------
print("\n[5/5] 归档位置")
arch = Path(moved[0]).parent if moved else None
check(arch is not None and arch.exists() and arch.name.startswith("重复截图_"),
      f"副本被移到了 {arch}（不是直接删除）")

print("\n结果:", "全部通过" if not BAD else f"{BAD} 项失败")
sys.exit(1 if BAD else 0)
