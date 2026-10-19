# -*- coding: utf-8 -*-
"""报告输出：Excel 结果表 + HTML 摘要

关键能力（2026-09-30 新增）：
- 截图落地：把查验截图整理成「原发票文件名_查验截图.png」，与原始发票
  一一对应，后续做账/装订时无需二次改名。
- HTML 内嵌缩略图（不再内嵌原图，报告体积从 8MB 降到 1MB 以内），
  并提供「下载全部查验截图」的 zip 打包入口。
"""
import base64
import datetime as _dt
import html
import io
import re
import shutil
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

STATUS_FILL = {
    "一致": "C6EFCE",      # 绿
    "冲红": "FFEB9C",      # 黄
    "作废": "FFEB9C",
    "不一致": "FFC7CE",    # 红
    "查无此票": "FFD9A0",   # 橙黄：需人工核对发票信息（不是查验通过！）
    "票面信息有误": "FFC7CE",  # 红：号码/日期/金额与平台校验不符
    "超次数跳过": "FFE0B2",  # 橙：当日次数用尽，明日重查
    "重复发票": "E4E4F7",   # 淡紫：同一张发票已被查验，结果继承
    "跳过（未识别为发票）": "EDEDED",
    "查验失败": "D9D9D9",  # 灰
    "平台限流": "F8CBAD",  # 橙红：被平台风控挡住，暂停后仍失败，需稍后重跑
    "未执行": "F2F2F2",    # 浅灰：本批被中断，这张还没查（可重跑）
}

STATUS_COLOR = {
    "一致": "#1a7f37", "冲红": "#b45309", "作废": "#b45309",
    "不一致": "#b91c1c", "查无此票": "#b45309",
    "票面信息有误": "#b91c1c", "超次数跳过": "#c2410c",
    "重复发票": "#5b4bb7", "跳过（未识别为发票）": "#6b7280",
    "查验失败": "#6b7280", "平台限流": "#c2410c", "未执行": "#6b7280",
}

COLS = ["序号", "文件名", "票种", "发票代码", "发票号码", "开票日期",
        "查验金额", "校验码", "查验结果", "明细", "查验截图文件"]

_SAFE_RE = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def safe_stem(name: str) -> str:
    """文件名净化：去掉 Windows 非法字符，限制长度"""
    s = _SAFE_RE.sub("_", str(name or "")).strip(" ._")
    return (s or "未命名")[:80]


# ---------------- Excel ----------------
def write_excel(results, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "查验结果"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center")
    disp = {r.get("idx"): i for i, r in enumerate(results, 1)}
    for i, r in enumerate(results, 1):
        dup = bool(r.get("is_dup") or r.get("dup_of"))
        shot_txt = r.get("shot_name", "") or (
            Path(r.get("shot", "")).name if r.get("shot") else "")
        if dup:
            # 重复行不落地截图，这里明确标注"共用第几行"，避免用户以为是漏截图
            shot_txt = (f"（与第 {disp.get(r.get('dup_of'), '?')} 行同一张票，"
                        f"共用同一张截图）")
        ws.append([i, Path(r["file"]).name, r.get("kind", ""),
                   r.get("code", ""), r.get("number", ""), r.get("date", ""),
                   r.get("amount", ""), r.get("checkcode", ""),
                   r.get("result", ""), r.get("detail", ""), shot_txt])
        row = ws[ws.max_row]
        fill = STATUS_FILL.get(r.get("result", ""))
        if fill:
            row[8].fill = PatternFill("solid", fgColor=fill)
        # 被去重合并的重复行：文件名列标淡紫，一眼看出"同一张票的其它副本"
        if r.get("is_dup"):
            row[1].fill = PatternFill("solid", fgColor="E4E4F7")
        # 查无此票是高优先级的人工复核项，明细整格标黄提醒
        if r.get("result") == "查无此票":
            row[9].fill = PatternFill("solid", fgColor="FFF2CC")
    widths = [6, 24, 12, 14, 22, 12, 10, 8, 12, 46, 34]
    for w, col in zip(widths, "ABCDEFGHIJK"):
        ws.column_dimensions[col].width = w
    wb.save(out_path)
    return out_path


# ---------------- 截图落地 / 打包 ----------------
def materialize_shots(results, out_dir, skip_dups=True) -> dict:
    """把查验截图整理成「原发票文件名_查验截图.png」，实现发票 ↔ 截图一一对应。

    - 普通行：直接复制/改名
    - 重复发票行（is_dup / dup_of）：**不再单独复制截图**（2026-10-08 改）。
      同一张发票只在截图目录里留一份，文件名取「主记录」——也就是导入清单中
      第一张不重复的那张发票的文件名。这样截图目录本身就是去重后的唯一集合，
      用户拿到后不必再手动删重复。
    返回 {记录 idx: 最终截图路径}（只含唯一发票的那一行）
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    mapping = {}
    for r in results:
        if skip_dups and (r.get("is_dup") or r.get("dup_of")):
            continue            # 重复副本不落地，避免截图目录出现一堆同图不同名的文件
        src = r.get("shot") or ""
        if not src or not Path(src).exists():
            continue
        base = safe_stem(Path(r.get("file") or r.get("number") or "发票").stem)
        dst = out / f"{base}_查验截图.png"
        k = 2
        while dst.exists() and Path(src).resolve() != dst.resolve():
            dst = out / f"{base}_查验截图_{k}.png"
            k += 1
        try:
            if Path(src).resolve() != dst.resolve():
                shutil.copy2(src, dst)
            mapping[r.get("idx")] = str(dst)
            r["shot_final"] = str(dst)
            r["shot_name"] = dst.name
        except Exception:
            continue
    return mapping


def prune_dup_shots(results, out_dir, archive_dir=None) -> list:
    """清掉历史遗留的「重复发票」截图副本（同一张票的其它文件名副本）。

    存在意义：早期版本会给重复行各复制一份截图，导致截图目录里堆了同图不同名的
    文件，用户还得手动去重。现在改为先把这些副本**移走**（不是删除），
    默认移到 历史输出/重复截图_<时间戳>/，确认无误后用户可自行删掉。

    判定很保守：只有「文件确实存在」且「内容与主记录截图逐字节相同」时才移走。
    返回被移走的文件**新路径**列表（都在 archive_dir 下）。
    """
    out = Path(out_dir)
    if not out.exists():
        return []
    by_idx = {r.get("idx"): r for r in results}
    moved = []
    targets = []
    for r in results:
        if not (r.get("is_dup") or r.get("dup_of")):
            continue
        m = by_idx.get(r.get("dup_of"))
        master = (m or {}).get("shot_final") or (m or {}).get("shot") or ""
        base = safe_stem(Path(r.get("file") or "").stem)
        cand = out / f"{base}_查验截图.png"
        if cand.exists() and master and Path(master).exists():
            targets.append((cand, master))
    if not targets:
        return []
    if archive_dir is None:
        archive_dir = out.parent / "历史输出" / (
            "重复截图_" + time.strftime("%Y%m%d_%H%M%S"))
    dst_dir = Path(archive_dir)
    dst_dir.mkdir(parents=True, exist_ok=True)
    for cand, master in targets:
        try:
            if _md5(cand) != _md5(master):
                continue        # 内容不一致 → 不是同一张票的副本，宁可不删
            newp = dst_dir / cand.name
            shutil.move(str(cand), str(newp))
            moved.append(str(newp))
        except Exception:
            continue
    return moved


def _md5(path) -> str:
    """文件指纹（截图只有几百 KB，直接整文件算，够快也够准）"""
    import hashlib
    try:
        h = hashlib.md5()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 256), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ""


def pack_shots_zip(mapping, zip_path, only_idx=None) -> str:
    """把整理好的截图打包成 zip（报告页「下载截图」用）。

    only_idx 给定时只打包这些记录 idx 的截图（用于"只下载查验成功的截图"）。
    """
    seen = set()
    try:
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
            for idx, p in mapping.items():
                if only_idx is not None and idx not in only_idx:
                    continue
                p = Path(p)
                if not p.exists():
                    continue
                arc, k = p.name, 2
                while arc in seen:
                    arc = f"{p.stem}_{k}{p.suffix}"
                    k += 1
                seen.add(arc)
                z.write(p, arc)
        return zip_path if seen else ""
    except Exception:
        return ""


# ---------------- HTML ----------------
def _thumb64(path, width=640) -> str:
    """缩略图 base64（JPEG，报告体积可控；原图仍保留在截图目录）"""
    try:
        from PIL import Image
        im = Image.open(path)
        if im.width > width:
            im = im.resize((width, max(1, int(im.height * width / im.width))),
                           Image.LANCZOS)
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=82)
        return base64.b64encode(buf.getvalue()).decode()
    except Exception:
        return ""


def write_html(results, out_path, shots_dir=None, zip_name=None,
               shots_out_dir=None, zip_ok_name=None):
    n_ok = sum(1 for r in results if r.get("result") == "一致")
    n_red = sum(1 for r in results if r.get("result") in ("冲红", "作废"))
    n_mis = sum(1 for r in results if r.get("result") == "不一致")
    n_none = sum(1 for r in results if r.get("result") == "查无此票")
    n_bad = sum(1 for r in results if r.get("result") in
                ("票面信息有误", "查验失败"))
    n_limit = sum(1 for r in results if r.get("result") == "超次数跳过")
    n_dup = sum(1 for r in results if r.get("is_dup"))
    n_skip = sum(1 for r in results if r.get("result") == "跳过（未识别为发票）")
    n_rate = sum(1 for r in results if r.get("result") == "平台限流")
    n_left = sum(1 for r in results if r.get("result") == "未执行")
    now = time.strftime("%Y-%m-%d %H:%M:%S")

    rows = []
    disp = {r.get("idx"): i for i, r in enumerate(results, 1)}
    for i, r in enumerate(results, 1):
        st = r.get("result", "")
        color = STATUS_COLOR.get(st, "#6b7280")
        dup = bool(r.get("is_dup") or r.get("dup_of"))
        if dup:
            # 重复副本不重复出图：只标注它共用哪一行的截图，报告体积也不再翻倍
            img = ""
            shot_name = (f"（与第 {disp.get(r.get('dup_of'), '?')} 行同一张票，"
                         f"共用同一张截图）")
        else:
            img = _thumb64(r.get("shot_final") or r.get("shot") or "")
            shot_name = r.get("shot_name", "")
        shot_name = html.escape(shot_name)
        rows.append(f"""
<tr><td>{i}</td><td>{html.escape(Path(r['file']).name)}</td>
<td>{html.escape(r.get('kind', ''))}</td>
<td>{html.escape(r.get('code', ''))}</td>
<td>{html.escape(r.get('number', ''))}</td>
<td>{html.escape(r.get('date', ''))}</td>
<td>{html.escape(r.get('amount', ''))}</td>
<td><b style="color:{color}">{html.escape(st)}</b></td>
<td>{html.escape(r.get('detail', ''))}</td>
<td style="font-size:11px;color:#6b7280">{shot_name}</td></tr>""")
        if img:
            rows.append(
                f'<tr class="shot"><td colspan="10">'
                f'<img loading="lazy" style="max-width:900px;border:1px solid '
                f'#ddd;border-radius:4px" src="data:image/jpeg;base64,{img}">'
                f"</td></tr>")

    zip_btn = ""
    # 张数一律按「唯一发票」算，与截图目录实际文件数一致
    n_shot = sum(1 for r in results
                 if (r.get("shot_final") or r.get("shot"))
                 and not (r.get("is_dup") or r.get("dup_of")))
    n_ok_shots = sum(1 for r in results if r.get("result") == "一致"
                     and not (r.get("is_dup") or r.get("dup_of")))
    if zip_ok_name:
        zip_btn += (f'<a class="btn" style="background:#1a7f37" '
                    f'href="{html.escape(zip_ok_name)}" download>'
                    f'⬇ 下载「查验一致」的截图（zip，{n_ok_shots} 张）</a> ')
    if zip_name:
        zip_btn += (f'<a class="btn" href="{html.escape(zip_name)}" download>'
                    f'⬇ 下载全部查验截图（zip，{n_shot} 张）</a>')
    dir_tip = ""
    if shots_out_dir:
        dir_tip = (f'<p class="tip">截图已按要求命名为 '
                   f'<b>「原发票文件名_查验截图.png」</b>，保存在：'
                   f'<code>{html.escape(str(shots_out_dir))}</code>'
                   f'（已去掉重复发票的副本，共 <b>{n_shot} 张</b>，'
                   f'可与原始发票直接一一对应使用，无需二次改名/去重）</p>')

    cards = [
        ("#1a7f37", f"查验一致 {n_ok}"),
        ("#b45309", f"冲红/作废 {n_red}"),
        ("#b91c1c", f"信息不一致 {n_mis}"),
        ("#b45309", f"查无此票（需核对）{n_none}"),
        ("#c2410c", f"超次数跳过 {n_limit}"),
        ("#5b4bb7", f"重复发票（已去重）{n_dup}"),
        ("#6b7280", f"票面有误/失败 {n_bad}"),
        ("#9ca3af", f"未识别跳过 {n_skip}"),
    ]
    if n_rate:
        cards.append(("#b91c1c", f"平台限流中断 {n_rate}"))
    if n_left:
        cards.append(("#6b7280", f"未执行（可重跑）{n_left}"))
    card_html = "".join(
        f'<div class="card" style="background:{bg}">{t}</div>' for bg, t in cards)

    notes = [
        "<b>查无此票</b>：平台未查到该发票 —— 通常是我们录入的号码/日期/金额有误，"
        "或该票信息本身有问题。<b>不代表查验通过</b>，请优先核对这几张。",
        "<b>超次数跳过</b>：该票今日查验次数已达平台上限（每票每天 5 次），"
        "程序已自动二次复核确认，明日重查即可。",
        "<b>重复发票（已去重）</b>：清单中同一张发票出现了多次，只提交查验一次，"
        "其余行继承同一结果；<b>截图目录里也只保留一份</b>"
        "（以导入清单中第一张不重复的发票文件名命名），无需手动去重。",
    ]
    if n_rate or n_left:
        notes.append(
            "<b>平台限流中断</b>：平台弹出「您的访问过于频繁，请稍后再试」，"
            "程序已自动暂停重试；若多次仍被限流则停止本批（"
            f"本批有 {n_rate} 张被限流、{n_left} 张未执行）。"
            "请等 10~30 分钟后用「导入Excel清单」只重跑这两类即可。")
    note_html = "• " + "<br>• ".join(notes)

    doc = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>发票查验结果 {now}</title>
<style>
body{{font-family:"Microsoft YaHei",sans-serif;margin:24px;color:#111;
background:#fff}}
h1{{font-size:20px;margin:0 0 6px}}
.sum{{display:flex;flex-wrap:wrap;gap:12px;margin:14px 0 10px}}
.card{{padding:10px 18px;border-radius:8px;font-size:14px;color:#fff}}
table{{border-collapse:collapse;width:100%;font-size:13px;margin-top:10px}}
td,th{{border:1px solid #ddd;padding:6px 10px;text-align:left}}
th{{background:#f3f4f6;position:sticky;top:0}}
tr:hover>td{{background:#f9fafb}}
tr.shot>td{{background:#fcfcfc;text-align:center;padding:10px}}
.btn{{display:inline-block;background:#1a4f8b;color:#fff;padding:10px 20px;
border-radius:6px;text-decoration:none;font-size:14px;margin:6px 0}}
.btn:hover{{background:#123a68}}
.tip{{color:#6b7280;font-size:12px;line-height:1.7}}
.note{{background:#fff8e6;border-left:4px solid #e6a23c;padding:10px 14px;
font-size:13px;line-height:1.8;border-radius:4px;margin:10px 0}}
</style></head><body>
<h1>发票查验结果汇总</h1>
<p class="tip">生成时间：{now}｜查验平台：inv-veri.chinatax.gov.cn｜
共 {len(results)} 条记录</p>
<div class="sum">
{card_html}
</div>
{zip_btn}
{dir_tip}
<div class="note">
<b>结果口径说明</b><br>
{note_html}
</div>
<table><tr><th>序号</th><th>文件名</th><th>票种</th><th>发票代码</th>
<th>发票号码</th><th>开票日期</th><th>金额</th><th>结果</th><th>明细</th>
<th>查验截图文件</th></tr>
{''.join(rows)}
</table></body></html>"""
    Path(out_path).write_text(doc, encoding="utf-8")
    return out_path


# ---------------- 清单读取 ----------------
def _clean_cell(v) -> str:
    """单元格 → 干净字符串（防 Excel 把长号码/日期转成数字）"""
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y%m%d%H%M%S")
    if isinstance(v, _dt.date):
        return v.strftime("%Y%m%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _norm_date(v) -> str:
    """开票日期 → YYYYMMDD（支持 2025-05-12 / 2025年5月12日 / datetime）"""
    s = _clean_cell(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return s[:8]
    digits = re.sub(r"\D", "", s)
    if len(digits) >= 8:
        return digits[:8]
    return digits


def read_scan_excel(path):
    """读取识别清单/查验结果 Excel → records 列表（导入功能）。
    表头模糊匹配：发票号码/发票代码/开票日期/金额/校验码/票种/文件名。
    发票号码为空的行标记为跳过项。"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    header = [str(c or "").strip() for c in rows[0]]

    def col(*kws):
        for i, h in enumerate(header):
            if any(k in h for k in kws):
                return i
        return None

    c_num = col("发票号码")
    c_code = col("发票代码")
    c_date = col("开票日期")
    c_amt = col("金额", "价税合计", "开具金额")
    c_chk = col("校验码")
    c_kind = col("票种", "kind")
    c_file = col("文件名", "文件")
    # 备选金额口径必须是「另一列」，避免与首选同列导致重复提交同一金额
    c_amt2 = None
    for i, h in enumerate(header):
        if i != c_amt and any(k in h for k in ("不含税", "未税", "不计税")):
            c_amt2 = i
            break
    out = []
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        num = _clean_cell(row[c_num]) if c_num is not None else ""
        num = re.sub(r"\.0$", "", num)
        rec = {
            "file": _clean_cell(row[c_file]) if c_file is not None else "Excel导入",
            "kind": _clean_cell(row[c_kind]) if c_kind is not None else "",
            "code": _clean_cell(row[c_code]) if c_code is not None else "",
            "number": num,
            "date": _norm_date(row[c_date]) if c_date is not None else "",
            "amount": (_clean_cell(row[c_amt]).replace("¥", "").replace("￥", "")
                       .replace(",", "").replace("，", "")
                       if c_amt is not None else ""),
            "amount_alt": (_clean_cell(row[c_amt2]).replace("¥", "")
                           .replace("￥", "").replace(",", "").replace("，", "")
                           if c_amt2 is not None else ""),
            "checkcode": _clean_cell(row[c_chk]) if c_chk is not None else "",
            "note": "",
        }
        if not rec["kind"]:
            rec["kind"] = "数电票" if len(rec["number"]) >= 18 else "增值税发票"
        if rec["amount_alt"] and rec["amount_alt"] == rec["amount"]:
            rec["amount_alt"] = ""
        out.append(rec)
    for i, rec in enumerate(out, 1):
        rec["idx"] = i
    return out
